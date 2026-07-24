"""
Buddha Travel and Tours — Sales Activity Log
A simple Streamlit form for staff to log Daily Summary and Enquiry entries.
Both write into a single Excel file (two sheets), safely handling multiple
staff submitting at the same time using a file lock.
"""

import streamlit as st
import openpyxl
from openpyxl import Workbook
from filelock import FileLock, Timeout
from datetime import date, datetime
import os

# ---------------------------------------------------------------------------
# CONFIG
# ---------------------------------------------------------------------------
EXCEL_PATH = "sales_activity_log.xlsx"
LOCK_PATH = EXCEL_PATH + ".lock"

DAILY_SUMMARY_HEADERS = [
    "Date",
    "Staff Name",
    "PNRs Generated Today",
    "Total Calls Received Today",
    "Total Emails Sent to Clients Today",
    "Total Enquiries Received & Logged in Zooma Today",
    "Total Callbacks Made Today (across all enquiries)",
    "Daily Tasks Planned / Noted for Today",
    "Tips/Tricks Applied for Price-Matching Objections",
    "Management Briefing Tips Applied Today?",
    "If Yes — Which Tips Applied",
    "Preferred Communication Method",
    "Client Feedback Received Today (if any)",
    "Services/Products Promoted Today",
    "Submitted At",
]

ENQUIRY_LOG_HEADERS = [
    "Date",
    "Staff Name",
    "Enquiry / Zooma Ref ID",
    "Client Name",
    "Date Enquiry Originally Received",
    "Enquiry Source / Channel",
    "Price/Options Offered?",
    "Price/Options Offered — Details",
    "Communication Method Used",
    "No. of Follow-Ups / Callbacks for This Enquiry",
    "Coordinated with a Teammate?",
    "Teammate Name (if Yes)",
    "Outcome",
    "Reason if Not Converted",
    "Notes",
    "Submitted At",
]

STAFF_NAMES = ["Please Select", "Kabita SYD", "Shreekantha KTM", "Prasant KTM", "Bishnu NP", "Nisha NP", "Bal Gopal MEL", "Prazol MEL", "Ram Hari NP", "Sabin NP", "Krishna SYD", "Srijana NP", "Uttam NP", "Mahendra NZ"]
COMMS_METHODS = ["WhatsApp", "SMS/Text", "Viber", "Phone Call", "Email"]
CHANNELS = ["Website Form", "Live Chat", "Facebook", "Instagram", "Phone", "Walk-in", "Referral", "Other"]
OUTCOMES = ["Quoted", "Booked (PNR generated)", "Lost", "Pending", "No Response"]
REASONS = [
    "Not Applicable / Converted",
    "Price / Lost on Fare",
    "Availability",
    "Client Still Deciding",
    "No Response from Client",
    "Went to Competitor",
    "Other",
]

# ---------------------------------------------------------------------------
# EXCEL HELPERS
# ---------------------------------------------------------------------------

def ensure_workbook_exists():
    """Create the workbook with both sheets and headers if it doesn't exist yet."""
    if not os.path.exists(EXCEL_PATH):
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Daily Summary"
        ws1.append(DAILY_SUMMARY_HEADERS)
        ws2 = wb.create_sheet("Enquiry Log")
        ws2.append(ENQUIRY_LOG_HEADERS)
        wb.save(EXCEL_PATH)


def append_row(sheet_name: str, row_values: list):
    """Safely append a row to the given sheet, using a file lock to avoid
    collisions if two staff submit at nearly the same moment."""
    lock = FileLock(LOCK_PATH, timeout=10)
    try:
        with lock:
            ensure_workbook_exists()
            wb = openpyxl.load_workbook(EXCEL_PATH)
            ws = wb[sheet_name]
            ws.append(row_values)
            wb.save(EXCEL_PATH)
        return True, None
    except Timeout:
        return False, "The file was busy — please try submitting again in a few seconds."
    except Exception as e:
        return False, f"Something went wrong saving your entry: {e}"


# ---------------------------------------------------------------------------
# STREAMLIT UI
# ---------------------------------------------------------------------------

st.set_page_config(page_title="Buddha Travel — Sales Activity Log", page_icon="✈️", layout="centered")

# Simple brand-colour header
st.markdown(
    """
    <div style="background-color:#0093d9; padding:16px 20px; border-radius:8px; margin-bottom:20px;">
        <h2 style="color:white; margin:0;">Buddha Travel and Tours</h2>
        <p style="color:#eaf6ff; margin:4px 0 0 0;">Sales Activity Log</p>
    </div>
    """,
    unsafe_allow_html=True,
)

entry_type = st.radio(
    "What are you submitting today?",
    ["Daily Summary", "Enquiry Log"],
    horizontal=True,
)

st.divider()

# ---------------------------------------------------------------------------
# DAILY SUMMARY FORM
# ---------------------------------------------------------------------------
if entry_type == "Daily Summary":
    with st.form("daily_summary_form", clear_on_submit=True):
        st.subheader("Daily Summary")

        col1, col2 = st.columns(2)
        with col1:
            entry_date = st.date_input("Date", value=date.today())
        with col2:
            staff_name = st.selectbox("Staff Name", STAFF_NAMES)

        pnrs = st.number_input("PNRs Generated Today", min_value=0, step=1)
        calls = st.number_input("Total Calls Received Today", min_value=0, step=1)
        emails = st.number_input("Total Emails Sent to Clients Today", min_value=0, step=1)
        enquiries_logged = st.number_input("Total Enquiries Received & Logged in Zooma Today", min_value=0, step=1)
        callbacks = st.number_input("Total Callbacks Made Today (across all enquiries)", min_value=0, step=1)

        tasks = st.text_area("Daily Tasks Planned / Noted for Today")
        price_tips = st.text_area("Tips/Tricks Applied for Price-Matching Objections (optional)")

        mgmt_tips_applied = st.radio("Management Briefing Tips Applied Today?", ["Yes", "No"], horizontal=True)
        mgmt_tips_detail = st.text_area("If Yes — Which Tips Applied (optional)")

        comms_pref = st.selectbox("Preferred Communication Method", COMMS_METHODS)
        feedback = st.text_area("Client Feedback Received Today (if any)")
        services_promoted = st.text_area("Services/Products Promoted Today")

        submitted = st.form_submit_button("Submit Daily Summary")

        if submitted:
            if staff_name == "Please select":
                st.error("Please select your staff name before submitting.")
            else:
                row = [
                    entry_date.strftime("%d/%m/%Y"),
                    staff_name,
                    int(pnrs),
                    int(calls),
                    int(emails),
                    int(enquiries_logged),
                    int(callbacks),
                    tasks,
                    price_tips,
                    mgmt_tips_applied,
                    mgmt_tips_detail,
                    comms_pref,
                    feedback,
                    services_promoted,
                    datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
                ]
                ok, err = append_row("Daily Summary", row)
                if ok:
                    st.success("Daily Summary submitted. Thank you!")
                else:
                    st.error(err)

# ---------------------------------------------------------------------------
# ENQUIRY LOG FORM
# ---------------------------------------------------------------------------
else:
    with st.form("enquiry_log_form", clear_on_submit=True):
        st.subheader("Enquiry Entry")

        col1, col2 = st.columns(2)
        with col1:
            entry_date = st.date_input("Date", value=date.today())
        with col2:
            staff_name = st.selectbox("Staff Name", STAFF_NAMES)

        ref_id = st.text_input("Enquiry / Zooma Ref ID")
        client_name = st.text_input("Client Name")
        date_received = st.date_input("Date Enquiry Originally Received", value=date.today())
        channel = st.selectbox("Enquiry Source / Channel", CHANNELS)

        price_offered = st.radio("Price/Options Offered?", ["Yes", "No"], horizontal=True)
        price_details = st.text_area("Price/Options Offered — Details (if Yes)")

        comms_used = st.selectbox("Communication Method Used", COMMS_METHODS)
        followups = st.number_input("No. of Follow-Ups / Callbacks for This Enquiry", min_value=0, step=1)

        coordinated = st.radio("Coordinated with a Teammate?", ["Yes", "No"], horizontal=True)
        teammate_name = st.selectbox("Teammate Name (if Yes)", STAFF_NAMES)

        outcome = st.selectbox("Outcome", OUTCOMES)
        reason = st.selectbox("Reason if Not Converted", REASONS)
        notes = st.text_area("Notes")

        submitted = st.form_submit_button("Submit Enquiry Entry")

        if submitted:
            if staff_name == "Please select":
                st.error("Please select your staff name before submitting.")
            elif not ref_id or not client_name:
                st.error("Please enter both the Enquiry/Zooma Ref ID and Client Name.")
            else:
                row = [
                    entry_date.strftime("%d/%m/%Y"),
                    staff_name,
                    ref_id,
                    client_name,
                    date_received.strftime("%d/%m/%Y"),
                    channel,
                    price_offered,
                    price_details,
                    comms_used,
                    int(followups),
                    coordinated,
                    teammate_name,
                    outcome,
                    reason,
                    notes,
                    datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
                ]
                ok, err = append_row("Enquiry Log", row)
                if ok:
                    st.success("Enquiry entry submitted. Thank you!")
                else:
                    st.error(err)

# ---------------------------------------------------------------------------
# ADMIN: DOWNLOAD CURRENT FILE (password protected)
# ---------------------------------------------------------------------------
st.divider()
with st.expander("Manager access: download current Excel file"):
    manager_password = st.text_input("Enter manager password", type="password")

    if manager_password:
        correct_password = st.secrets.get("manager_password", None)
        if correct_password is None:
            st.error("No manager password has been set up yet. Add one in Streamlit Cloud → Settings → Secrets.")
        elif manager_password == correct_password:
            ensure_workbook_exists()
            with open(EXCEL_PATH, "rb") as f:
                st.download_button(
                    label="Download sales_activity_log.xlsx",
                    data=f,
                    file_name="sales_activity_log.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
        else:
            st.error("Incorrect password.")
